"""
WO-42 — bulk payments that Finance can actually pay, and an auditor can
actually read.

Three gaps closed, each test named after the failure it prevents:
  * the payee schedule Finance pays from (full bank details, finance/admin
    only, every download on the audit trail, only once approved);
  * the audit packet that listed no payees on a 100-person payment;
  * voucher pages 2..N that did not say which voucher they belonged to.

Run: python test_payee_schedule.py
"""
from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="docex-sched-")
os.environ["DOCEX_ORG"] = "sched"

import store  # noqa: E402

store.set_store(store.JsonFileStore(_TMP))

import auth as A  # noqa: E402

A._SECRET_FILE = Path(_TMP) / ".s"
A._secret_cache = None

import org_config  # noqa: E402
import payment_voucher as pv  # noqa: E402
import requisition_export  # noqa: E402
import requisitions as rq  # noqa: E402

ORG = "sched"
org_config.set_features(ORG, multi_payee_requisitions=True)
PW = "correct-horse-battery"
TEMPLATE = {
    "title": "Payment Voucher",
    "letterhead": {"org_name": "Test Foundation", "rc_number": "1"},
    "signature_roles": [{"label": "Prepared by", "source": "submitter"}],
    "pv_number_format": "TF/{MONYY}/PV/{seq}",
}

_passed = _failed = 0


def check(label: str, cond: bool, detail: str = "") -> None:
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  ok   {label}")
    else:
        _failed += 1
        print(f"  FAIL {label}{(' — ' + detail) if detail else ''}")


def _payees(n: int) -> list[rq.Payee]:
    return [rq.Payee(name=f"Participant {i}", account_number=f"01234{i:05d}",
                     bank_name="GTBank" if i % 2 else "Zenith Bank",
                     amount=15000.0 + i, purpose="Transport refund, Maiduguri training")
            for i in range(1, n + 1)]


def _bulk(n: int = 100, *, status: str | None = None, org: str = ORG) -> rq.Requisition:
    total = sum(p.amount for p in _payees(n))
    req = rq.create_requisition(org, submitted_by="officer@t", department="program",
                                vendor_name="Workshop participants", amount=total,
                                payees=_payees(n))
    if status:
        # Reaching APPROVED through the real chain is covered by the
        # requisition suites; here only the state matters.
        raw = store.get_store().get(org, "requisitions", req.id)
        raw["status"] = status
        store.get_store().put(org, "requisitions", req.id, raw)
        req = rq.get_requisition(org, req.id)
    return req


def _xlsx_rows(content: bytes) -> list[list]:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(content)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _pdf_pages(content: bytes) -> list[str]:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return [(p.extract_text() or "") for p in pdf.pages]


# ─── the schedule, engine level ─────────────────────────────────────────────


def test_finance_never_pays_from_a_schedule_that_disagrees_with_the_voucher() -> None:
    print("\nThe schedule and the voucher are the same payment, to the kobo")
    pv.set_template(ORG, TEMPLATE)
    req = _bulk(100, status="approved")
    content = pv.payee_schedule_xlsx(req, ORG, generated_by="fin@t")
    rows = _xlsx_rows(content)
    text = "\n".join(" ".join(str(c) for c in r if c is not None) for r in rows)
    voucher = pv.build_voucher(req, ORG)

    payee_rows = [r for r in rows if r and str(r[1] or "").startswith("Participant ")]
    check("one row per payee", len(payee_rows) == 100, str(len(payee_rows)))
    check("full account numbers — this is the file Finance pays from",
          "0123400100" in text and "0123400001" in text)
    check("each row carries its own bank", "Zenith Bank" in text and "GTBank" in text)
    total_row = [r for r in rows if r and r[1] == "TOTAL"]
    check("there is a TOTAL row", bool(total_row))
    check("schedule total == voucher total",
          bool(total_row) and abs(float(total_row[0][4]) - voucher.total) < 0.005,
          f"{total_row[0][4] if total_row else None} vs {voucher.total}")
    check("the schedule names its voucher", voucher.pv_number and voucher.pv_number in text,
          voucher.pv_number)
    check("and its requisition", req.ref in text)
    check("and says who produced it", "fin@t" in text)
    check("and warns that it carries bank details", "CONFIDENTIAL" in text.upper())


def test_a_schedule_that_does_not_add_up_is_refused_not_printed() -> None:
    print("\nIf the payee lines no longer sum to the approved amount, no schedule")
    pv.set_template(ORG, TEMPLATE)
    req = _bulk(3, status="approved")
    raw = store.get_store().get(ORG, "requisitions", req.id)
    raw["amount"] = raw["amount"] + 1000          # someone edited the amount alone
    store.get_store().put(ORG, "requisitions", req.id, raw)
    req = rq.get_requisition(ORG, req.id)
    try:
        pv.payee_schedule_xlsx(req, ORG, generated_by="fin@t")
        check("refused", False, "a schedule was produced")
    except pv.ScheduleError as exc:
        check("refused, and the reason names both figures",
              "1,000" in str(exc) or "differ" in str(exc).lower(), str(exc))


def test_a_single_payee_payment_still_gets_a_one_line_schedule() -> None:
    print("\nA one-vendor payment gets a one-line schedule from the vendor fields")
    pv.set_template(ORG, TEMPLATE)
    req = rq.create_requisition(ORG, submitted_by="officer@t", department="program",
                                vendor_name="Acme Supplies Ltd", amount=250000.0,
                                vendor_account="0987654321", vendor_bank_name="Access Bank")
    raw = store.get_store().get(ORG, "requisitions", req.id)
    raw["status"] = "paid"
    store.get_store().put(ORG, "requisitions", req.id, raw)
    req = rq.get_requisition(ORG, req.id)
    text = str(_xlsx_rows(pv.payee_schedule_xlsx(req, ORG, generated_by="fin@t")))
    check("vendor, bank and full account on one line",
          "Acme Supplies Ltd" in text and "Access Bank" in text and "0987654321" in text)


def test_nobody_pays_from_a_schedule_before_the_payment_is_approved() -> None:
    print("\nA schedule exists only once every approver has signed")
    for status, allowed in [("draft", False), ("submitted", False), ("in_review", False),
                            ("on_hold", False), ("returned", False), ("declined", False),
                            ("approved", True), ("paid", True)]:
        req = _bulk(2)
        req.status = status
        check(f"{status:9s} ⇒ {'allowed' if allowed else 'refused'}",
              (pv.schedule_refusal(req) is None) == allowed, str(pv.schedule_refusal(req)))


def test_who_may_download_bank_details_is_the_org_s_choice() -> None:
    print("\nFinance by default; an org that calls it something else can say so")
    check("default: finance department", pv.schedule_departments(ORG) == ["finance"],
          str(pv.schedule_departments(ORG)))
    other = "sched_other"
    pv.set_template(other, {**TEMPLATE, "schedule_departments": ["accounts", "treasury"]})
    check("configured: that org's own departments",
          pv.schedule_departments(other) == ["accounts", "treasury"],
          str(pv.schedule_departments(other)))
    check("and it does not leak into the first org", pv.schedule_departments(ORG) == ["finance"])


# ─── the schedule, through the real API ─────────────────────────────────────


def test_bank_details_leave_the_system_only_through_the_right_hands() -> None:
    print("\nHTTP: flag, role, status and the audit trail, all enforced server-side")
    from fastapi.testclient import TestClient
    import api.main as m

    c = TestClient(m.app, raise_server_exceptions=False)
    A.create_user("admin@t", "Admin", PW, "management", "admin", org_id=ORG)
    A.create_user("fin@t", "Finance Officer", PW, "finance", "approver", org_id=ORG)
    A.create_user("prog@t", "Programme Officer", PW, "program", "approver", org_id=ORG)

    def headers(email):
        tok = c.post("/auth/login", json={"email": email, "password": PW}).json()["token"]
        return {"Authorization": f"Bearer {tok}"}

    admin, fin, prog = headers("admin@t"), headers("fin@t"), headers("prog@t")
    pv.set_template(ORG, TEMPLATE)
    req = _bulk(5, status="approved")
    url = f"/requisitions/{req.ref}/payees.xlsx"
    before = len(rq.get_requisition(ORG, req.id).audit_log)

    org_config.set_features(ORG, payee_schedule_export=False, voucher_export=True)
    perms = c.get("/requisitions/document-permissions", headers=fin).json()
    check("flag OFF ⇒ the screen is told not to show the schedule button",
          perms.get("payee_schedule") is False and perms.get("voucher") is True, str(perms))
    r = c.get(url, headers=fin)
    check("flag OFF ⇒ 404, as if the route did not exist", r.status_code == 404, str(r.status_code))

    org_config.set_features(ORG, payee_schedule_export=True)
    for who, h, expect in (("programme officer", prog, False), ("finance officer", fin, True),
                           ("admin", admin, True)):
        perms = c.get("/requisitions/document-permissions", headers=h).json()
        check(f"the screen offers the schedule to a {who}: {expect}",
              perms.get("payee_schedule") is expect, str(perms))
    r = c.get(url, headers=prog)
    check("a programme officer is refused (403)", r.status_code == 403, str(r.status_code))
    check("and a refused download leaves no trace in the trail",
          len(rq.get_requisition(ORG, req.id).audit_log) == before)

    r = c.get(url, headers=fin)
    check("a finance officer gets the file", r.status_code == 200, f"{r.status_code} {r.text[:120]}")
    check("it is a spreadsheet", r.content[:2] == b"PK")
    check("named for the PV number Finance files by",
          "TF-" in r.headers.get("content-disposition", ""), r.headers.get("content-disposition", ""))

    after = rq.get_requisition(ORG, req.id)
    last = after.audit_log[-1]
    check("the download is on the audit trail", last.event == "payee_schedule_downloaded", last.event)
    check("naming who downloaded it", last.actor == "fin@t", last.actor)
    check("how many people and how much", "5 payees" in last.detail, last.detail)
    check("and the chain still verifies", rq.verify_audit_chain(after))

    r = c.get(url, headers=admin)
    check("an admin gets the file too", r.status_code == 200, str(r.status_code))
    check("and that download is recorded separately",
          rq.get_requisition(ORG, req.id).audit_log[-1].actor == "admin@t")

    pending = _bulk(2, status="in_review")
    r = c.get(f"/requisitions/{pending.ref}/payees.xlsx", headers=fin)
    check("not yet approved ⇒ 409, and says why", r.status_code == 409 and "approv" in r.text.lower(),
          f"{r.status_code} {r.text[:120]}")


# ─── the audit packet ────────────────────────────────────────────────────────


def test_the_auditor_sees_every_payee_but_not_their_full_account_numbers() -> None:
    print("\nThe audit packet lists every payee, accounts masked to the last four")
    req = _bulk(100)
    pages = _pdf_pages(requisition_export.requisition_pdf(req))
    text = "\n".join(pages)
    check("the first payee is listed", "Participant 1 " in text or "Participant 1\n" in text)
    check("the hundredth payee is listed", "Participant 100" in text)
    check("last four digits are shown — enough to match a bank statement",
          "0100" in text and "0001" in text)
    check("no full account number appears anywhere",
          "0123400100" not in text and "0123400001" not in text)
    check("the payee total is stated", "1,505,050" in text, "expected 1,505,050.00")

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(requisition_export.requisition_xlsx(req)))
    check("the Excel version has a Payees sheet", "Payees" in wb.sheetnames, str(wb.sheetnames))
    if "Payees" in wb.sheetnames:
        flat = str([[c.value for c in r] for r in wb["Payees"].iter_rows()])
        check("with all 100 payees", "Participant 100" in flat)
        check("and masked accounts there too", "0123400100" not in flat and "0100" in flat)


def test_a_single_payee_packet_is_unchanged() -> None:
    print("\nA one-vendor packet still shows vendor detail, not an empty payee table")
    req = rq.create_requisition(ORG, submitted_by="o@t", department="program",
                                vendor_name="Acme Supplies Ltd", amount=1000.0,
                                vendor_account="0987654321", vendor_bank_name="Access Bank")
    text = "\n".join(_pdf_pages(requisition_export.requisition_pdf(req)))
    check("vendor detail present", "Access Bank" in text)
    check("no payee table", "Payees (" not in text)


# ─── the voucher's pages ─────────────────────────────────────────────────────


def test_a_page_that_falls_out_of_the_file_still_says_whose_it_is() -> None:
    print("\nEvery voucher page carries the PV number and 'page X of Y'")
    pv.set_template(ORG, TEMPLATE)
    req = _bulk(100, status="approved")
    pages = _pdf_pages(pv.voucher_pdf(req, ORG))
    number = pv.pv_number_for(req, ORG)
    n = len(pages)
    check("a 100-payee voucher runs over several pages", n >= 3, str(n))
    for i, page in enumerate(pages, start=1):
        check(f"page {i} names {number} and says {i} of {n}",
              number in page and f"Page {i} of {n}" in page, page[-160:].replace("\n", " | "))
    short = pv.voucher_pdf(_bulk(1, status="approved"), ORG)
    check("a one-page voucher says Page 1 of 1", "Page 1 of 1" in _pdf_pages(short)[0])


def test_applying_a_profile_installs_its_voucher_template() -> None:
    print("\nThe voucher in the profile is the voucher the client gets")
    import json
    prof = json.loads(Path("profiles/neem.json").read_text())
    prof = {**prof, "org_id": "applied"}
    prof["admin"] = {**prof["admin"], "password": "Passw0rd!Passw0rd!"}
    org_config.apply_profile(prof)
    tpl = pv.get_template("applied")
    want = prof["voucher_template"]
    check("letterhead org name comes from the profile",
          (tpl.get("letterhead") or {}).get("org_name") == want["letterhead"]["org_name"],
          str(tpl.get("letterhead")))
    check("PV number format comes from the profile",
          tpl.get("pv_number_format") == want.get("pv_number_format"), str(tpl.get("pv_number_format")))
    bad = {**prof, "org_id": "applied_bad", "voucher_template": "not-a-template"}
    rep = org_config.validate_profile(bad)
    check("a malformed voucher_template is rejected BEFORE anything is written",
          not rep.ok and any("voucher_template" in e for e in rep.errors), str(rep.errors))
    check("and nothing was written for that org",
          not store.get_store().get("applied_bad", "config", "features"))


def test_the_browser_can_read_the_filename_the_server_chose() -> None:
    print("\nCross-site downloads keep the PV-number filename")
    from fastapi.testclient import TestClient
    import api.main as m
    c = TestClient(m.app, raise_server_exceptions=False)
    r = c.options("/health", headers={"Origin": "http://localhost:3000",
                                      "Access-Control-Request-Method": "GET"})
    r = c.get("/health", headers={"Origin": "http://localhost:3000"})
    exposed = r.headers.get("access-control-expose-headers", "").lower()
    check("Content-Disposition is exposed to the browser", "content-disposition" in exposed, exposed)


if __name__ == "__main__":
    print("WO-42 — bulk payments: schedule, audit packet, voucher pages")
    test_finance_never_pays_from_a_schedule_that_disagrees_with_the_voucher()
    test_a_schedule_that_does_not_add_up_is_refused_not_printed()
    test_a_single_payee_payment_still_gets_a_one_line_schedule()
    test_nobody_pays_from_a_schedule_before_the_payment_is_approved()
    test_who_may_download_bank_details_is_the_org_s_choice()
    test_bank_details_leave_the_system_only_through_the_right_hands()
    test_the_auditor_sees_every_payee_but_not_their_full_account_numbers()
    test_a_single_payee_packet_is_unchanged()
    test_a_page_that_falls_out_of_the_file_still_says_whose_it_is()
    test_applying_a_profile_installs_its_voucher_template()
    test_the_browser_can_read_the_filename_the_server_chose()
    print(f"\n{_passed} passed, {_failed} failed")
    raise SystemExit(1 if _failed else 0)
