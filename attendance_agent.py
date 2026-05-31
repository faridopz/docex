"""
DOCex Attendance Payment Agent.

Targets the NGO programs-team pain point of paying event attendees: today
the team manually cross-references TWO files — the attendance log (who
showed up each day) and the payment info form (who's registered + their
bank details) — then multiplies days × rate, builds a payment schedule,
thumb-types every account number into a bank app to verify, and emails
the schedule to finance. The whole process takes 3-4 hours per event
for a 30-person workshop.

This agent automates steps 1-5 of that flow:

  1. Parse the attendance log (name + per-day attendance marks)
  2. Parse the payment info form (name + org + bank + account)
  3. Fuzzy cross-match names across both files using the same blended
     scorer Bank Verify uses (WRatio + token_sort_ratio, processor=str.lower)
  4. Bucket each person: paid / no_attendance / no_payment_info
  5. Calculate days_attended × rate_per_day for the paid bucket

Step 6 (run the schedule through Bank Verify) and step 7 (email finance)
are handled by the API layer that composes this engine with the existing
Bank Verify primitive — keeping each piece replaceable.

This is the first *composite* agent — it doesn't introduce a new core
capability, it chains existing ones. That choice is deliberate: every
primitive we build (parse, match, verify, email) is now reusable in the
next agent (Sub-award Agent, Procurement Agent, etc.) without duplication.
"""
from __future__ import annotations

import logging
from datetime import datetime, timezone
from typing import Optional

from openpyxl import load_workbook
from rapidfuzz import fuzz

from bank_verify import resolve_bank_code
from models import (
    AttendancePaymentRun,
    AttendanceRecord,
    MatchedAttendee,
    PaymentInfoRecord,
    RateCard,
)

logger = logging.getLogger(__name__)


# ─── Column detection ───────────────────────────────────────────────────────
#
# Both files arrive with column headers chosen by the team that prepared
# them — there's no canonical schema. We auto-detect by keyword, same
# pattern as Bank Verify's parse_schedule. When we can't find a header
# row, we raise ValueError with a message the UI surfaces to the user.

_HEADER_SCAN_DEPTH = 15

# Attendance log columns — at minimum we need a name column. Day columns
# can be detected by date-looking headers, "Day 1/2/3" style headers, or
# weekday names. Anything else is ignored.
_ATTENDANCE_NAME_KEYWORDS = (
    "name", "participant", "attendee", "delegate", "recipient",
)
_DAY_KEYWORDS = (
    "day ", "day-", "day1", "day2", "day3", "day4", "day5",
    "monday", "tuesday", "wednesday", "thursday", "friday",
    "saturday", "sunday",
    "mon ", "tue ", "wed ", "thu ", "fri ", "sat ", "sun ",
)

# Payment info form columns — need name + account + bank at minimum.
# Org is optional but useful when present (lets the schedule group by
# partner org for ED sign-off, future iteration).
_PAY_NAME_KEYWORDS = (
    "name", "participant", "recipient", "payee", "beneficiary",
)
_PAY_ORG_KEYWORDS = (
    "organisation", "organization", "org", "agency", "partner", "company",
    "employer", "institution",
)
_PAY_ACCOUNT_KEYWORDS = ("account", "acct", "nuban", "a/c")
_PAY_BANK_KEYWORDS = ("bank",)
# Role detection — used to pick the right rate from a RateCard. Common
# Nigerian NGO payment-form labels include "Role", "Position",
# "Designation", "Title", "Category".
_PAY_ROLE_KEYWORDS = ("role", "position", "designation", "title", "category", "function")


def _find_attendance_columns(sheet) -> tuple[int, int, list[tuple[int, str]]]:
    """
    Locate the header row in an attendance log.

    Returns (header_row_index, name_col, [(day_col, day_label), ...]).
    Day columns are kept ordered left-to-right so day numbering on the
    output ("Day 1", "Day 2") matches the file's natural reading order.

    Raises ValueError when no row contains a name column AND ≥1 day column.
    """
    max_col = sheet.max_column or 1

    for row_idx in range(1, min(_HEADER_SCAN_DEPTH, sheet.max_row or 1) + 1):
        name_col: Optional[int] = None
        day_cols: list[tuple[int, str]] = []
        for col in range(1, max_col + 1):
            raw = sheet.cell(row=row_idx, column=col).value
            if raw is None:
                continue
            text = str(raw).strip()
            if not text:
                continue
            lower = text.lower()
            # Name column — match once, first hit wins
            if name_col is None and any(k in lower for k in _ATTENDANCE_NAME_KEYWORDS):
                name_col = col
                continue
            # Day column — accept anything that looks like a date or day marker
            if _looks_like_day(lower, raw):
                day_cols.append((col, text))

        if name_col is not None and day_cols:
            return row_idx, name_col, day_cols

    raise ValueError(
        "Could not find a header row with a name column and at least one day "
        "column in the attendance log. The day columns should be labelled "
        "with dates (e.g. '2026-05-25') or day numbers ('Day 1', 'Day 2'). "
        "Rename your headers and try again."
    )


def _looks_like_day(lower: str, raw) -> bool:
    """Heuristic: does this header value represent a day in the event?"""
    # Explicit keyword match (Day 1, Mon, Monday, etc.)
    if any(k in lower for k in _DAY_KEYWORDS):
        return True
    # Date-typed cells from Excel — openpyxl returns datetime objects
    if isinstance(raw, datetime):
        return True
    # Looks like a date string (yyyy-mm-dd, dd/mm/yyyy, dd-mm-yyyy)
    if any(c.isdigit() for c in lower):
        digit_count = sum(1 for c in lower if c.isdigit())
        if digit_count >= 4 and any(sep in lower for sep in ("-", "/", ".")):
            return True
    return False


def _find_payment_info_columns(sheet) -> tuple[int, dict[str, int]]:
    """
    Locate the header row in a payment info form.

    Returns (header_row_index, {"name": col, "account": col, "bank": col,
    "org": col_or_None}).
    """
    max_col = sheet.max_column or 1

    for row_idx in range(1, min(_HEADER_SCAN_DEPTH, sheet.max_row or 1) + 1):
        col_map: dict[str, int] = {}
        for col in range(1, max_col + 1):
            raw = sheet.cell(row=row_idx, column=col).value
            if raw is None:
                continue
            text = str(raw).lower().strip()
            if not text:
                continue
            # Order matters — bank should not be classified as name etc.
            if "bank" in text and "code" not in text and "bank" not in col_map:
                col_map["bank"] = col
                continue
            if any(k in text for k in _PAY_ACCOUNT_KEYWORDS) and "account" not in col_map:
                col_map["account"] = col
                continue
            if any(k in text for k in _PAY_ORG_KEYWORDS) and "org" not in col_map:
                col_map["org"] = col
                continue
            if any(k in text for k in _PAY_NAME_KEYWORDS) and "name" not in col_map:
                col_map["name"] = col
                continue
            if any(k in text for k in _PAY_ROLE_KEYWORDS) and "role" not in col_map:
                col_map["role"] = col
                continue

        if {"name", "account", "bank"}.issubset(col_map):
            return row_idx, col_map

    raise ValueError(
        "Could not find a header row with 'name', 'account', and 'bank' "
        "columns in the payment info form. Rename your column headers "
        "(e.g. 'Recipient Name', 'Account Number', 'Bank') and try again."
    )


# ─── Parsing ────────────────────────────────────────────────────────────────


def _is_attendance_mark_present(value) -> bool:
    """
    A cell counts as "present" if it's clearly affirmative. We accept the
    range of marks teams actually use on attendance sheets:
      - "x", "X", "✓", "yes", "y", "p", "present", "1", "T", "TRUE"
    Empty strings, "no", "n", "0", "absent" all count as not-present.
    """
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    s = str(value).strip().lower()
    if not s:
        return False
    return s in {
        "x", "✓", "✔", "yes", "y", "p", "present", "1", "true", "t",
        "attended", "here", "shown",
    }


def parse_attendance_log(file_path_or_buffer) -> list[AttendanceRecord]:
    """
    Parse an attendance log .xlsx into a list of AttendanceRecord.

    The file is expected to have one row per attendee and one column per
    day of the event, plus a name column. Day cells should be marked with
    one of the standard affirmative marks (x, ✓, yes, p, present, 1, etc.)
    when the person attended that day.

    Returns one record per attendee, with the count + labels of days they
    were present. Attendees with zero days attended are still returned —
    the downstream matcher needs them to populate the no_attendance bucket.
    """
    wb = load_workbook(file_path_or_buffer, data_only=True)
    sheet = wb.active

    header_row, name_col, day_cols = _find_attendance_columns(sheet)

    records: list[AttendanceRecord] = []
    for excel_row in range(header_row + 1, (sheet.max_row or header_row) + 1):
        name_value = sheet.cell(row=excel_row, column=name_col).value
        if not name_value:
            continue
        name = str(name_value).strip()
        if not name:
            continue

        day_labels: list[str] = []
        for col, label in day_cols:
            cell_value = sheet.cell(row=excel_row, column=col).value
            if _is_attendance_mark_present(cell_value):
                day_labels.append(label)

        records.append(
            AttendanceRecord(
                name=name,
                days_attended=len(day_labels),
                day_labels=day_labels,
            )
        )

    return records


def parse_payment_info(file_path_or_buffer) -> list[PaymentInfoRecord]:
    """Parse the payment info form .xlsx into PaymentInfoRecord rows."""
    wb = load_workbook(file_path_or_buffer, data_only=True)
    sheet = wb.active

    header_row, col_map = _find_payment_info_columns(sheet)

    records: list[PaymentInfoRecord] = []
    for excel_row in range(header_row + 1, (sheet.max_row or header_row) + 1):
        name_v = sheet.cell(row=excel_row, column=col_map["name"]).value
        account_v = sheet.cell(row=excel_row, column=col_map["account"]).value
        bank_v = sheet.cell(row=excel_row, column=col_map["bank"]).value
        if not (name_v and account_v and bank_v):
            continue

        # Account number can come through as a float ("7042310445.0") — strip
        account_str = str(account_v).replace(" ", "").replace("-", "").strip()
        if account_str.endswith(".0"):
            account_str = account_str[:-2]

        bank_str = str(bank_v).strip()
        # resolve_bank_code handles "Opay" → "999992", "GTB" → "058", etc.
        # Pass through unknowns so the downstream Bank Verify call can
        # surface a clean error rather than us silently dropping the row.
        bank_code = resolve_bank_code(bank_str) or bank_str

        org_v = None
        if "org" in col_map:
            raw_org = sheet.cell(row=excel_row, column=col_map["org"]).value
            if raw_org:
                org_v = str(raw_org).strip()

        role_v = None
        if "role" in col_map:
            raw_role = sheet.cell(row=excel_row, column=col_map["role"]).value
            if raw_role:
                role_v = str(raw_role).strip()

        records.append(
            PaymentInfoRecord(
                name=str(name_v).strip(),
                organisation=org_v,
                account_number=account_str,
                bank_code=bank_code,
                bank_name=bank_str,
                role=role_v,
            )
        )

    return records


# ─── Matching ───────────────────────────────────────────────────────────────
#
# Same blended scorer as Bank Verify — WRatio (catches within-token typos)
# clamped by token_sort_ratio (catches shared-prefix attacks). Processor
# str.lower normalises case across the two files.

_MATCH_VERIFIED_THRESHOLD = 85
_MATCH_WARNING_THRESHOLD = 65


def _name_match_score(a: str, b: str) -> int:
    w = fuzz.WRatio(a, b, processor=str.lower)
    t = fuzz.token_sort_ratio(a, b, processor=str.lower)
    return int(min(w, t + 10))


# ─── Rate-card application ──────────────────────────────────────────────────


def build_rate_card_from_flat(rate_per_day: float) -> RateCard:
    """
    Wrap a single flat rate in a degenerate RateCard so the engine has
    one rate-application code path regardless of whether the user supplied
    a card or a flat number. The flat-rate run wizard creates one of these
    on the fly; the rate-card picker passes a real saved card.
    """
    return RateCard(
        id="ad-hoc",
        name=f"Flat ₦{rate_per_day:,.0f}/day",
        default_rate_per_day=rate_per_day,
        roles=[],
    )


def _resolve_rate(role: Optional[str], card: RateCard) -> float:
    """Look up a role on the card; fall back to the card's default rate."""
    if role:
        role_lower = role.strip().lower()
        for line in card.roles:
            if line.role.strip().lower() == role_lower:
                return float(line.amount_per_day)
    return float(card.default_rate_per_day)


# ─── Accuracy flags ─────────────────────────────────────────────────────────
#
# Generated AFTER matching. Each flag is a human-readable string surfaced
# to the user above the verify button so the team's existing "double-check
# before payment" ritual becomes an explicit step in the product rather
# than a manual habit. The Bank Verify primitive is still the final gate
# at payment time — these flags catch issues earlier, when they're cheaper.


def _detect_accuracy_flags(
    payment_info: list[PaymentInfoRecord],
    matched: list[MatchedAttendee],
    no_attendance: list[MatchedAttendee],
    no_payment_info: list[MatchedAttendee],
    days_in_event: int,
) -> list[str]:
    flags: list[str] = []

    # Duplicate account numbers — same NUBAN appears on multiple rows of
    # the payment info form. Usually a data entry mistake; sometimes
    # legitimate (one person registered twice, or two people share an
    # account). Either way worth surfacing.
    by_account: dict[str, list[str]] = {}
    for p in payment_info:
        by_account.setdefault(p.account_number, []).append(p.name)
    for acct, names in by_account.items():
        if len(names) > 1:
            flags.append(
                f"Duplicate account {acct} on payment info — appears for "
                f"{len(names)} people: {', '.join(names[:3])}"
                + (f" (+{len(names) - 3} more)" if len(names) > 3 else "")
            )

    # Implausible attendance — someone marked present on more days than
    # the event actually ran. Suggests duplicate day columns or a row that
    # got marked across the entire grid.
    for m in matched:
        if days_in_event > 0 and m.days_attended > days_in_event:
            flags.append(
                f"{m.payment_info_name or m.attendance_name}: marked "
                f"{m.days_attended} days on a {days_in_event}-day event"
            )

    # Low-confidence matches in the paid bucket — score below the verified
    # threshold. We still pay them (the team chose to release at this score)
    # but flag for human eyeball check before money moves.
    low_score = [m for m in matched if (m.match_score or 0) < 85]
    if low_score:
        flags.append(
            f"{len(low_score)} paid {'row' if len(low_score) == 1 else 'rows'} "
            f"matched with confidence below 85 — confirm identity: "
            + ", ".join(
                f"{m.payment_info_name} (score {m.match_score})"
                for m in low_score[:3]
            )
            + (f" (+{len(low_score) - 3} more)" if len(low_score) > 3 else "")
        )

    # No-payment-info attendees in significant numbers — suggests the
    # payment info form is incomplete or stale.
    if len(no_payment_info) >= 3:
        flags.append(
            f"{len(no_payment_info)} attendees missing bank info — chase "
            f"them for payment details before this batch goes out"
        )

    # No-attendance count proportional to total — if more than half the
    # payment info form didn't actually attend, something's off with the
    # attendance file (wrong event? wrong date range?).
    total_registered = len(payment_info)
    if total_registered > 0 and len(no_attendance) > total_registered * 0.5:
        flags.append(
            f"{len(no_attendance)}/{total_registered} registered people "
            f"did NOT attend — double-check the attendance file is for "
            f"the right event"
        )

    return flags


# ─── Google Sheets ingestion ────────────────────────────────────────────────
#
# NGOs collect event registrations through Google Forms. Forms feed a
# Google Sheet 1:1. So if we can accept a Google Sheet URL, we accept
# Forms responses with zero extra work. Same for attendance — many teams
# tick attendance on a shared Google Sheet too.
#
# Implementation: extract the sheet id from the share URL, fetch the
# /export?format=xlsx URL, hand the bytes to the existing openpyxl parser.
# Cleaner than reinventing the parser for CSV: same column detection, same
# day-mark recognition, same robustness against title rows.


def google_sheets_xlsx_bytes(url: str, timeout_seconds: float = 30.0) -> bytes:
    """
    Convert a Google Sheets share URL into the raw bytes of its .xlsx
    export. The sheet must be 'Anyone with the link can view' (no auth
    flow yet — that's a Phase 2 OAuth project).

    Accepts any of these URL shapes:
      https://docs.google.com/spreadsheets/d/{ID}/edit?usp=sharing
      https://docs.google.com/spreadsheets/d/{ID}/edit#gid=0
      https://docs.google.com/spreadsheets/d/{ID}
    Raises ValueError if the URL doesn't look like a Google Sheet, or
    httpx.HTTPError on network/auth failures.
    """
    import re

    import httpx

    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url)
    if not match:
        raise ValueError(
            "URL doesn't look like a Google Sheets share link. Expected "
            "something like https://docs.google.com/spreadsheets/d/.../edit"
        )
    sheet_id = match.group(1)
    export_url = (
        f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    )

    resp = httpx.get(export_url, timeout=timeout_seconds, follow_redirects=True)
    # Google returns 200 with an HTML "you need to sign in" page when a
    # private sheet is requested without auth. Detect by content-type.
    ctype = resp.headers.get("content-type", "")
    if resp.status_code != 200 or "text/html" in ctype:
        raise ValueError(
            "Could not fetch the Google Sheet. Make sure share access is "
            "set to 'Anyone with the link can view' — DOCex can't read "
            "private sheets yet."
        )
    return resp.content


def match_and_build_run(
    event_name: str,
    rate_per_day: float,
    attendance: list[AttendanceRecord],
    payment_info: list[PaymentInfoRecord],
    rate_card: Optional[RateCard] = None,
) -> AttendancePaymentRun:
    """
    Cross-match attendance ↔ payment info and produce an AttendancePaymentRun.

    Matching algorithm:
      For each payment_info record:
        Find the best-scoring attendance record above the warning threshold.
        If found:
          - Pair them; calc amount = days × rate.
          - Bucket as "paid" if days_attended ≥ 1, else "no_attendance".
        If not found:
          - Bucket as "no_attendance" (registered, didn't show).

      For each attendance record not matched to any payment_info:
        - Bucket as "no_payment_info" (showed up, can't pay them).

    Greedy 1:1 matching — once an attendance record is matched to one
    payee it can't be matched again. This handles the rare case where
    two payees have similar names (e.g. "Amina Hassan" + "Amina Mohammed
    Hassan") by giving priority to the higher-scoring pair.

    Rates are applied via the RateCard. If no card is supplied, we wrap
    the flat rate_per_day in a degenerate single-rate card so the engine
    has one code path. When a payee's role column matches a role on the
    card, that role's rate is applied; otherwise the card's default rate
    is applied.
    """
    # Normalise to a card. Flat-rate callers get a one-line degenerate card.
    if rate_card is None:
        rate_card = build_rate_card_from_flat(rate_per_day)

    days_in_event = max(
        (r.days_attended for r in attendance),
        default=0,
    )
    # Better signal of event length: pull from union of day_labels across
    # all attendees. Someone might have attended only day 1 of a 3-day event.
    all_day_labels: set[str] = set()
    for r in attendance:
        for label in r.day_labels:
            all_day_labels.add(label)
    if all_day_labels:
        days_in_event = len(all_day_labels)

    # Greedy 1:1 matching. Sort candidate pairs by score descending and
    # claim each (payee, attendee) at most once.
    candidate_pairs: list[tuple[int, int, int]] = []  # (score, pay_idx, att_idx)
    for pi, pay in enumerate(payment_info):
        for ai, att in enumerate(attendance):
            score = _name_match_score(pay.name, att.name)
            if score >= _MATCH_WARNING_THRESHOLD:
                candidate_pairs.append((score, pi, ai))
    candidate_pairs.sort(reverse=True)

    matched_pay_idx: dict[int, tuple[int, int]] = {}  # pay_idx → (att_idx, score)
    claimed_att: set[int] = set()
    for score, pi, ai in candidate_pairs:
        if pi in matched_pay_idx:
            continue
        if ai in claimed_att:
            continue
        matched_pay_idx[pi] = (ai, score)
        claimed_att.add(ai)

    paid: list[MatchedAttendee] = []
    no_attendance: list[MatchedAttendee] = []

    for pi, pay in enumerate(payment_info):
        pair = matched_pay_idx.get(pi)
        # Resolve the rate once per payee — same look-up regardless of
        # which bucket they end up in. Carrying applied_rate_per_day on
        # the result lets the audit say "₦45k = 3 days × ₦15k/day,
        # Facilitator rate from card 'TA Connect 2026'".
        applied_rate = _resolve_rate(pay.role, rate_card)

        if pair is None:
            # Registered but no attendance row matched at all → never showed
            no_attendance.append(
                MatchedAttendee(
                    payment_info_name=pay.name,
                    attendance_name=None,
                    match_score=None,
                    status="no_attendance",
                    days_attended=0,
                    day_labels=[],
                    organisation=pay.organisation,
                    account_number=pay.account_number,
                    bank_code=pay.bank_code,
                    bank_name=pay.bank_name,
                    role=pay.role,
                    applied_rate_per_day=applied_rate,
                    amount=0.0,
                )
            )
            continue

        att_idx, score = pair
        att = attendance[att_idx]
        # Even if we matched, the person may have attended 0 days (matched
        # name but every day cell empty). That belongs in no_attendance too.
        if att.days_attended == 0:
            no_attendance.append(
                MatchedAttendee(
                    payment_info_name=pay.name,
                    attendance_name=att.name,
                    match_score=score,
                    status="no_attendance",
                    days_attended=0,
                    day_labels=[],
                    organisation=pay.organisation,
                    account_number=pay.account_number,
                    bank_code=pay.bank_code,
                    bank_name=pay.bank_name,
                    role=pay.role,
                    applied_rate_per_day=applied_rate,
                    amount=0.0,
                )
            )
            continue

        paid.append(
            MatchedAttendee(
                payment_info_name=pay.name,
                attendance_name=att.name,
                match_score=score,
                status="paid",
                days_attended=att.days_attended,
                day_labels=att.day_labels,
                organisation=pay.organisation,
                account_number=pay.account_number,
                bank_code=pay.bank_code,
                bank_name=pay.bank_name,
                role=pay.role,
                applied_rate_per_day=applied_rate,
                amount=att.days_attended * applied_rate,
            )
        )

    no_payment_info: list[MatchedAttendee] = []
    for ai, att in enumerate(attendance):
        if ai in claimed_att:
            continue
        # An attendee with zero days attended who never matched a payee
        # is a no-op (probably a blank row). Skip to keep the bucket
        # focused on people who actually need follow-up.
        if att.days_attended == 0:
            continue
        no_payment_info.append(
            MatchedAttendee(
                payment_info_name=None,
                attendance_name=att.name,
                match_score=None,
                status="no_payment_info",
                days_attended=att.days_attended,
                day_labels=att.day_labels,
                organisation=None,
                account_number=None,
                bank_code=None,
                bank_name=None,
                amount=0.0,
            )
        )

    total_to_pay = sum(r.amount for r in paid)

    # Generate the accuracy panel — the user reviews this BEFORE Bank
    # Verify fires. Bank Verify remains the final hard gate at payment
    # time; these flags are the earlier, cheaper soft gate.
    accuracy_flags = _detect_accuracy_flags(
        payment_info=payment_info,
        matched=paid,
        no_attendance=no_attendance,
        no_payment_info=no_payment_info,
        days_in_event=days_in_event,
    )

    return AttendancePaymentRun(
        event_name=event_name,
        rate_per_day=rate_card.default_rate_per_day,
        days_in_event=days_in_event,
        matched=paid,
        no_attendance=no_attendance,
        no_payment_info=no_payment_info,
        total_to_pay=total_to_pay,
        paid_count=len(paid),
        no_attendance_count=len(no_attendance),
        no_payment_info_count=len(no_payment_info),
        created_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        rate_card_id=rate_card.id if rate_card.id != "ad-hoc" else None,
        rate_card_name=rate_card.name,
        rate_card_snapshot=rate_card,
        accuracy_flags=accuracy_flags,
    )
