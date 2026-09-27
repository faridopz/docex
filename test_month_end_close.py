"""
Close the month: reconcile, explain, sign off, hand QuickBooks one file.

What was wrong with month end, found by running a realistic statement:
  * every bank charge — SMS alert, stamp duty, NIP fee, VAT on the fee — was a
    HIGH "money left with no approval" finding. A normal Nigerian statement
    carries dozens, so the real finding drowned in noise;
  * nothing proved the statement was complete (opening + in − out = closing),
    and money coming IN was ignored entirely;
  * one person closed the month; nobody reviewed it;
  * the QuickBooks "payment register" had no import screen to go to in the
    international edition of QuickBooks Online (the one Nigerian organisations
    use): Intuit's own staff confirm it cannot import journal entries or
    expenses. Bank transactions (Date, Description, Amount) are the path every
    edition has.

Run: python test_month_end_close.py
"""
from __future__ import annotations

import csv
import io
import os
import tempfile
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="docex-close-")
os.environ["DOCEX_ORG"] = "acme"

import store  # noqa: E402

store.set_store(store.JsonFileStore(_TMP))

import auth as A  # noqa: E402

A._SECRET_FILE = Path(_TMP) / ".s"
A._secret_cache = None

import accounting_export as ax  # noqa: E402
import bank_accounts as ba  # noqa: E402
import bank_reconciliation as br  # noqa: E402
import org_config  # noqa: E402
import requisitions as rq  # noqa: E402

_passed = _failed = 0
ORG = "acme"
PERIOD = "2026-09"
PW = "correct-horse-battery"


def check(label: str, cond: bool, detail: str = "") -> None:
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  ok   {label}")
    else:
        _failed += 1
        print(f"  FAIL {label}{(' — ' + detail) if detail else ''}")


# Opening balance 1,000,000.00. Running balances are exact.
ROWS = [
    ("01/09/2026", "CARE FCDO GRANT TRANCHE Q3", "INW001", "", "5,000,000.00"),
    ("03/09/2026", "TRF TO HOTEL MAIDUGURI LTD", "FT26090301", "850,000.00", ""),
    ("03/09/2026", "NIP TRANSFER CHARGE", "", "26.88", ""),
    ("03/09/2026", "VAT ON NIP CHARGE", "", "2.02", ""),
    ("05/09/2026", "BULK PAYMENT WORKSHOP STIPENDS", "BULK0905", "600,000.00", ""),
    ("10/09/2026", "STAMP DUTY", "", "50.00", ""),
    ("15/09/2026", "SMS ALERT CHARGES AUG", "", "4.00", ""),
    ("20/09/2026", "TRF TO JOHN OKAFOR", "FT26092077", "120,000.00", ""),
    ("25/09/2026", "LOAN ARRANGEMENT CHARGE", "", "450,000.00", ""),
    ("30/09/2026", "ACCOUNT MAINTENANCE FEE", "", "1,075.00", ""),
]


def _statement(rows=ROWS, opening=1_000_000.00) -> bytes:
    bal = opening
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Trans Date", "Narration", "Reference", "Debit", "Credit", "Balance"])
    for d, narr, ref, debit, credit in rows:
        amt = float(credit.replace(",", "") or 0) - float(debit.replace(",", "") or 0)
        bal = round(bal + amt, 2)
        w.writerow([d, narr, ref, debit, credit, f"{bal:,.2f}"])
    return buf.getvalue().encode()


def _approved(amount, vendor, category, project="B24", payees=None):
    r = rq.create_requisition(ORG, submitted_by="raiser@acme.org", department="program",
                              vendor_name=vendor, amount=amount, category=category,
                              project_code=project, payees=payees)
    raw = store.get_store().get(ORG, "requisitions", r.id)
    raw["status"] = "approved"
    store.get_store().put(ORG, "requisitions", r.id, raw)
    return r


def _setup():
    org_config.set_features(ORG, bank_reconciliation=True, multi_payee_requisitions=True,
                            accounting_export=True)
    wf = rq.get_workflow(ORG)
    wf.allowed_categories = []
    rq.set_workflow(ORG, wf)
    acct = ba.add(ORG, code="B24", name="CARE / FCDO", account_number="0011223344",
                  bank_name="GTBank", project_code="B24")
    ba.set_quickbooks_name(ORG, acct.id, "GTBank CARE 3344")
    ax.set_map(ORG, ax.AccountMap(
        accounts={"venue": "Venue costs", "workshop": "Participant stipends"},
        classes={"B24": "CARE FCDO"},
        bank_charges_account="Bank charges",
        edition="online_international"))
    hotel = _approved(850_000, "Hotel Maiduguri", "venue")
    rq.mark_paid(ORG, hotel.id, actor="fin@acme.org", bank_reference="FT26090301",
                 account_id=acct.id)
    payees = [rq.Payee(name=f"Participant {i}", account_number=f"01234{i:05d}",
                       bank_name="GTBank", amount=15_000.0) for i in range(40)]
    batch = _approved(600_000, "Workshop", "workshop", payees=payees)
    rq.mark_paid(ORG, batch.id, actor="fin@acme.org", bank_reference="BULK0905",
                 account_id=acct.id, settlement="bulk")
    for email, dept, role in (("fin@acme.org", "finance", "approver"),
                              ("tunde@acme.org", "finance", "approver"),
                              ("prog@acme.org", "program", "reviewer")):
        A.create_user(email, email.split("@")[0], PW, dept, role, org_id=ORG)
    return acct, hotel, batch


def _run(acct, rows=ROWS):
    return br.reconcile(ORG, PERIOD, _statement(rows), filename="gtb.csv",
                        account_id=acct.id, actor="fin@acme.org")


def _codes(run):
    out = {}
    for e in run.exceptions:
        out.setdefault(e.code.value, []).append(e)
    return out


def test_charges_are_recognised_not_flagged(acct) -> None:
    print("\nBank charges are grouped and totalled — not 'money left with no approval'")
    run = _run(acct)
    codes = _codes(run)
    charges = codes.get("BANK_CHARGE", [])
    check("five small charges recognised", len(charges) == 5,
          str([e.description for e in charges]))
    check("none of them is a high-severity finding",
          all(e.severity == br.Severity.LOW for e in charges))
    s = br.summary(run)
    check("charges totalled", abs(s["bank_charges_total"] - 1_157.90) < 0.005, str(s.get("bank_charges_total")))
    high = [e for e in run.exceptions if e.severity == br.Severity.HIGH]
    check("the real findings stay high: the unknown transfer AND the large 'charge'",
          sorted(e.description for e in high) == ["LOAN ARRANGEMENT CHARGE", "TRF TO JOHN OKAFOR"],
          str([e.description for e in high]))


def test_both_approved_payments_match(acct) -> None:
    print("\nThe hotel payment and the 40-person bulk schedule each match their one debit")
    run = _run(acct)
    refs = sorted(m.vendor_name for m in run.matches)
    check("two matches", len(run.matches) == 2, str(refs))
    check("including the bulk schedule as one debit", any("40 payees" in r for r in refs), str(refs))


def test_money_in_and_completeness(acct) -> None:
    print("\nMoney in is listed; the balance column proves the statement is complete")
    run = _run(acct)
    s = br.summary(run)
    check("money in totalled", abs(s["money_in_total"] - 5_000_000) < 0.005, str(s.get("money_in_total")))
    check("opening and closing balances read",
          abs(s["opening_balance"] - 1_000_000) < 0.005 and abs(s["closing_balance"] - 3_978_842.10) < 0.005,
          f"{s.get('opening_balance')} / {s.get('closing_balance')}")
    check("statement complete", s["statement_complete"] is True, str(s.get("balance_gaps")))

    gappy = [r for r in ROWS if r[1] != "STAMP DUTY"]
    full = _statement()
    # Rebuild the statement with its balances, then remove one row: the
    # balances no longer chain, which is exactly what a missing page looks like.
    lines = full.decode().splitlines()
    removed = "\n".join(l for l in lines if "STAMP DUTY" not in l).encode()
    run2 = br.reconcile(ORG, PERIOD, removed, filename="gtb.csv", account_id=acct.id, actor="fin@acme.org")
    s2 = br.summary(run2)
    check("a missing row is detected", s2["statement_complete"] is False and s2["balance_gaps"],
          str(s2.get("balance_gaps")))
    check("and the gap is flagged as a finding",
          any(e.code.value == "STATEMENT_GAP" for e in run2.exceptions))


def test_explain_many_close_and_review(acct) -> None:
    print("\nExplain several at once; close; a DIFFERENT person signs it off")
    run = _run(acct)
    high = [e for e in run.exceptions if e.severity == br.Severity.HIGH]
    run = br.explain_many(ORG, run.id, bank_line_ids=[e.bank_line_id for e in high],
                          actor="fin@acme.org", reason="Loan fee agreed with bank; John Okafor refund to donor")
    check("both explained in one step", run.unresolved_high == 0, str(run.unresolved_high))
    run = br.close_period(ORG, run.id, actor="fin@acme.org")
    try:
        br.review_period(ORG, run.id, actor="fin@acme.org")
        check("closer cannot sign off their own month", False)
    except br.ReconciliationError as exc:
        check("closer cannot sign off their own month", "different" in str(exc).lower(), str(exc))
    run = br.review_period(ORG, run.id, actor="tunde@acme.org")
    check("signed off by the second person", run.reviewed_by == "tunde@acme.org" and run.reviewed_at)
    return run


def test_quickbooks_bank_file(run) -> None:
    print("\nQuickBooks bank file: every line, plain numbers, and each line says what it is")
    files = ax.bank_upload_from_run(ORG, run)
    rows = list(csv.reader(io.StringIO(files[0])))
    check("3 columns, QuickBooks' own names", rows[0] == ["Date", "Description", "Amount"], str(rows[0]))
    check("every statement line in the month", len(rows) - 1 == len(ROWS), str(len(rows) - 1))
    hotel = next(r for r in rows if "Hotel Maiduguri" in r[1])
    check("matched line names payee, request, account and class",
          all(x in hotel[1] for x in ("REQ-", "Venue costs", "CARE FCDO")), hotel[1])
    check("amount is a bare signed number, date day-first", hotel[2] == "-850000.00" and hotel[0] == "03/09/2026",
          str(hotel))
    bulk = next(r for r in rows if "Participant stipends" in r[1])
    check("the bulk schedule is one line, coded", "40 payees" in bulk[1] and bulk[2] == "-600000.00", str(bulk))
    charge = next(r for r in rows if "STAMP DUTY" in r[1])
    check("charges are labelled with the charges account", "Bank charges" in charge[1], charge[1])
    grant = next(r for r in rows if "GRANT TRANCHE" in r[1])
    check("money in is included, positive", grant[2] == "5000000.00", str(grant))
    check("no ₦ signs or thousands commas anywhere", "₦" not in files[0] and "5,000" not in files[0])


def test_journals_only_from_a_closed_month(acct, run) -> None:
    print("\nCoded journal (QuickBooks US / Desktop) — only from a closed month, and balanced")
    open_run = _run(acct)
    try:
        ax.journal_csv_from_run(ORG, open_run)
        check("refused while the month is open", False)
    except ax.ExportError as exc:
        check("refused while the month is open", "close" in str(exc).lower(), str(exc))
    text = ax.journal_csv_from_run(ORG, run)
    rows = list(csv.DictReader(io.StringIO(text)))
    check("QuickBooks' journal columns",
          set(["Journal No.", "Journal Date", "Account Name", "Debits", "Credits", "Description", "Class"])
          <= set(rows[0].keys()), str(list(rows[0].keys())))
    by_no = {}
    for r in rows:
        by_no.setdefault(r["Journal No."], []).append(r)
    balanced = all(abs(sum(float(x["Debits"] or 0) for x in v) - sum(float(x["Credits"] or 0) for x in v)) < 0.005
                   for v in by_no.values())
    check("every journal balances", balanced)
    hotel = next(v for v in by_no.values() if any("Hotel Maiduguri" in x["Description"] for x in v))
    check("hotel: debit Venue costs (class CARE FCDO), credit the GTBank account",
          any(x["Account Name"] == "Venue costs" and x["Debits"] == "850000.00" and x["Class"] == "CARE FCDO" for x in hotel)
          and any(x["Account Name"] == "GTBank CARE 3344" and x["Credits"] == "850000.00" for x in hotel),
          str(hotel))
    check("bank charges posted as one journal", sum(1 for k in by_no if "CHG" in k) == 1, str(list(by_no)))
    check("unexplained money is never posted", not any("JOHN OKAFOR" in r["Description"].upper() for r in rows))
    again = ax.journal_csv_from_run(ORG, run)
    check("re-downloading gives the same journal numbers (QuickBooks flags a double import)",
          [r["Journal No."] for r in csv.DictReader(io.StringIO(again))] == [r["Journal No."] for r in rows])
    iif = ax.journal_iif_from_run(ORG, run)
    check("Desktop IIF: GENERAL JOURNAL transactions that sum to zero",
          iif.startswith("!TRNS") and "GENERAL JOURNAL" in iif and _iif_balanced(iif), iif[:200])


def _iif_balanced(iif: str) -> bool:
    total, ok = 0.0, True
    for line in iif.splitlines():
        parts = line.split("\t")
        if parts[0] in ("TRNS", "SPL"):
            total += float(parts[6])
        elif parts[0] == "ENDTRNS":
            ok = ok and abs(total) < 0.005
            total = 0.0
    return ok


def test_report_for_auditors(run) -> None:
    print("\nReconciliation report (Excel) for auditors and donors")
    data = ax.reconciliation_report_xlsx(ORG, run)
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    check("is an Excel workbook", any(n.startswith("xl/worksheets/") for n in names))
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    check("summary, matched, attention, charges and money-in sheets",
          {"Summary", "Matched", "Explained & open items", "Bank charges", "Money in"} <= set(wb.sheetnames),
          str(wb.sheetnames))
    text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value is not None)
    check("names who closed and who signed off", "fin@acme.org" in text and "tunde@acme.org" in text)
    check("states the statement was complete", "complete" in text.lower())


def test_http(acct, hotel) -> None:
    print("\nHTTP: the month-end screen's calls")
    from fastapi.testclient import TestClient
    import api.main as m
    c = TestClient(m.app, raise_server_exceptions=False)

    def h(email):
        return {"Authorization": "Bearer " + c.post("/auth/login", json={"email": email, "password": PW}).json()["token"]}

    fin, tunde, prog = h("fin@acme.org"), h("tunde@acme.org"), h("prog@acme.org")
    r = c.post("/reconciliation", headers=fin, data={"period": PERIOD, "account_id": acct.id},
               files={"statement": ("gtb.csv", _statement(), "text/csv")})
    check("reconciled", r.status_code == 200, f"{r.status_code} {r.text[:160]}")
    run = r.json()
    rid = run["id"]
    check("the screen gets the charges and completeness", run["summary"]["bank_charges_count"] == 5
          and run["summary"]["statement_complete"] is True, str(run.get("summary"))[:300])
    r = c.get(f"/reconciliation/{rid}/quickbooks", headers=fin, params={"kind": "bank"})
    check("bank file downloads", r.status_code == 200 and r.text.startswith("Date,Description,Amount"), r.text[:80])
    check("named for the account and month", "2026-09" in r.headers.get("content-disposition", ""),
          r.headers.get("content-disposition", ""))
    high = [e["bank_line_id"] for e in run["exceptions"] if e["severity"] == "high"]
    r = c.post(f"/reconciliation/{rid}/explain-many", headers=fin,
               json={"bank_line_ids": high, "reason": "checked with the bank"})
    check("explain several", r.status_code == 200 and r.json()["summary"]["unresolved_high"] == 0,
          f"{r.status_code} {r.text[:160]}")
    r = c.get(f"/reconciliation/{rid}/quickbooks", headers=fin, params={"kind": "journal"})
    check("journal refused while open (409)", r.status_code == 409, f"{r.status_code} {r.text[:120]}")
    c.post(f"/reconciliation/{rid}/close", headers=fin)
    r = c.post(f"/reconciliation/{rid}/review", headers=fin)
    check("closer cannot review own (400)", r.status_code == 400, str(r.status_code))
    r = c.post(f"/reconciliation/{rid}/review", headers=tunde)
    check("second person signs off", r.status_code == 200 and r.json()["reviewed_by"] == "tunde@acme.org",
          f"{r.status_code} {r.text[:160]}")
    r = c.get(f"/reconciliation/{rid}/quickbooks", headers=fin, params={"kind": "journal"})
    check("journal downloads once closed", r.status_code == 200 and "Journal No." in r.text, f"{r.status_code}")
    r = c.get(f"/reconciliation/{rid}/report.xlsx", headers=fin)
    check("report downloads", r.status_code == 200 and r.content[:2] == b"PK", str(r.status_code))
    r = c.get(f"/reconciliation/{rid}/quickbooks", headers=prog, params={"kind": "bank"})
    check("staff outside Finance cannot take the books out (403)", r.status_code == 403, str(r.status_code))


if __name__ == "__main__":
    print("Close the month — reconcile, explain, sign off, one file for QuickBooks")
    acct, hotel, batch = _setup()
    test_charges_are_recognised_not_flagged(acct)
    test_both_approved_payments_match(acct)
    test_money_in_and_completeness(acct)
    closed = test_explain_many_close_and_review(acct)
    test_quickbooks_bank_file(closed)
    test_journals_only_from_a_closed_month(acct, closed)
    test_report_for_auditors(closed)
    test_http(acct, hotel)
    print(f"\n{_passed} passed, {_failed} failed")
    raise SystemExit(1 if _failed else 0)
