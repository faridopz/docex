"""
End-to-end HTTP test: WO-25, per-requisition PDF/Excel export + the
weekly/monthly audit-log export.

NEEM's ask, from the original feedback dump: "per-requisition PDF/Excel
export + monthly/weekly log export for audit". Reuses the exact
StreamingResponse/Content-Disposition pattern already proven twice
(api/bank_verify_routes.py, api/attendance_agent_routes.py) and the
per-requisition Attachment download route's raw-Response pattern for the
PDF. requisition_export.py (the pure builder module) never touches
FastAPI or the network, so its output can be verified directly — this
suite checks the actual returned bytes are a real PDF (%PDF header,
readable by pdfplumber, contains the requisition's ref/vendor/comment
text) and a real .xlsx (readable by openpyxl, right sheet names, right
cell values) rather than just asserting a 200 status.

Covers: the flag gates all three endpoints; a per-requisition export
contains the requisition's real data (not a fixed template); the log
export is bounded correctly by date (a requisition outside the range is
excluded, one inside is included); an empty range 404s with a clear
message; a malformed date 400s; end-before-start is refused; visibility
matches viewing (no extra role gate beyond being signed in — any
authenticated user who could GET the requisition can export it).

Run: python test_requisition_export.py
"""
from __future__ import annotations

import io
import tempfile
from pathlib import Path

import store
import auth as auth_mod
import org_config

_base = Path(tempfile.mkdtemp(prefix="docex_export_"))
store.set_store(store.JsonFileStore(_base / "store"))
auth_mod._SECRET_FILE = _base / ".auth_secret"
auth_mod._secret_cache = None

from fastapi.testclient import TestClient  # noqa: E402
import api.main as m  # noqa: E402

client = TestClient(m.app)
_fail = 0


def check(name, cond):
    global _fail
    print(f"{'PASS' if cond else 'FAIL'}  {name}")
    if not cond:
        _fail += 1


def hdr(email, password):
    tok = client.post("/auth/login", json={"email": email, "password": password}).json()["token"]
    return {"Authorization": f"Bearer {tok}"}


# ─── bootstrap ──────────────────────────────────────────────────────────────

client.post("/auth/register", json={
    "email": "admin@neem.org", "name": "Admin", "password": "admin-passphrase",
    "department": "management", "role": "admin"})
ah = hdr("admin@neem.org", "admin-passphrase")

client.post("/auth/register", headers=ah, json={
    "email": "program@neem.org", "name": "Program", "password": "program-passphrase",
    "department": "program", "role": "reviewer"})
ph = hdr("program@neem.org", "program-passphrase")

r = client.put("/requisitions/workflow", headers=ah, json={
    "steps": [{"key": "finance", "label": "Finance", "department": "finance"}],
})
check("workflow saved", r.status_code == 200)

import json as _json

r = client.post("/requisitions", headers=ph, data={
    "vendor_name": "Greenview Caterers", "amount": "75000", "category": "catering",
    "vendor_bank_name": "Zenith Bank", "vendor_tin": "TIN-55210",
    "vendor_phone_or_email": "billing@greenview.example",
    "payment_type": "advance",
    "budget_lines": _json.dumps([
        {"description": "Workshop catering", "unit": "Person", "budget_line": "B-300",
         "quantity": 25, "frequency": 1, "unit_cost": 3000},
    ]),
})
check("requisition raised", r.status_code == 200)
req_id, req_ref = r.json()["id"], r.json()["ref"]
check("budget line total is server-computed (25 * 1 * 3000)",
      r.json()["budget_lines"][0]["line_total"] == 75000)
client.post(f"/requisitions/{req_id}/comments", headers=ph, data={"text": "Vendor confirmed for the workshop."})

# ─── flag off: all three refused ────────────────────────────────────────────

r = client.get(f"/requisitions/{req_id}/export.pdf", headers=ph)
check("flag off: PDF export refused", r.status_code == 400)
check("refusal names the feature flag", "requisition_export" in r.json()["detail"])

r = client.get(f"/requisitions/{req_id}/export.xlsx", headers=ph)
check("flag off: xlsx export refused", r.status_code == 400)

r = client.get("/requisitions/export/log.xlsx", headers=ph, params={"start": "2020-01-01", "end": "2030-01-01"})
check("flag off: log export refused", r.status_code == 400)

org_config.set_features("default", requisition_export=True)

# ─── per-requisition PDF: real content, not a fixed template ───────────────

r = client.get(f"/requisitions/{req_id}/export.pdf", headers=ph)
check("PDF export succeeds", r.status_code == 200)
check("PDF content-type", r.headers.get("content-type", "").startswith("application/pdf"))
check("filename in Content-Disposition", req_ref in r.headers.get("content-disposition", ""))
check("looks like a real PDF", r.content.startswith(b"%PDF"))

import pdfplumber
with pdfplumber.open(io.BytesIO(r.content)) as pdf:
    pdf_text = "\n".join((p.extract_text() or "") for p in pdf.pages)
check("PDF contains the requisition ref", req_ref in pdf_text)
check("PDF contains the vendor name", "Greenview Caterers" in pdf_text)
check("PDF contains the comment text", "Vendor confirmed for the workshop" in pdf_text)
check("PDF contains the submitter", "program@neem.org" in pdf_text)
check("PDF confirms the audit chain verified", "Hash chain verified" in pdf_text)
check("PDF shows the payment type", "Advance" in pdf_text)
check("PDF shows the amount in words", "Seventy-Five Thousand Naira Only" in pdf_text)
check("PDF shows the vendor's bank", "Zenith Bank" in pdf_text)
check("PDF shows the vendor's TIN", "TIN-55210" in pdf_text)
check("PDF shows the budget-line description", "Workshop catering" in pdf_text)
check("PDF shows the computed budget-line total", "75,000.00" in pdf_text)

# ─── per-requisition Excel: real workbook, real sheets, real values ────────

r = client.get(f"/requisitions/{req_id}/export.xlsx", headers=ph)
check("xlsx export succeeds", r.status_code == 200)
check("xlsx content-type", "spreadsheetml" in r.headers.get("content-type", ""))
check("looks like a real xlsx (zip header)", r.content.startswith(b"PK"))

from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(r.content))
check("Summary sheet present", "Summary" in wb.sheetnames)
check("Comments sheet present", "Comments" in wb.sheetnames)
check("Audit Log sheet present", "Audit Log" in wb.sheetnames)
summary_values = [c.value for row in wb["Summary"].iter_rows() for c in row]
check("Summary sheet carries the ref", req_ref in summary_values)
check("Summary sheet carries the vendor", "Greenview Caterers" in summary_values)
check("Summary sheet carries the amount in words",
      "Seventy-Five Thousand Naira Only" in summary_values)
check("Summary sheet carries the vendor's bank", "Zenith Bank" in summary_values)
check("Budget Lines sheet present", "Budget Lines" in wb.sheetnames)
budget_values = [c.value for row in wb["Budget Lines"].iter_rows() for c in row]
check("Budget Lines sheet carries the description", "Workshop catering" in budget_values)
check("Budget Lines sheet carries the server-computed total", 75000.0 in budget_values)

# ─── nonexistent requisition 404s on both formats ──────────────────────────

r = client.get("/requisitions/does-not-exist/export.pdf", headers=ph)
check("nonexistent requisition PDF 404s", r.status_code == 404)
r = client.get("/requisitions/does-not-exist/export.xlsx", headers=ph)
check("nonexistent requisition xlsx 404s", r.status_code == 404)

# ─── log export: date-bounded correctly ────────────────────────────────────

import datetime as dt
today = dt.date.today()
today_s = today.isoformat()
yesterday_s = (today - dt.timedelta(days=1)).isoformat()
tomorrow_s = (today + dt.timedelta(days=1)).isoformat()
far_past_start = (today - dt.timedelta(days=400)).isoformat()
far_past_end = (today - dt.timedelta(days=399)).isoformat()

# `dt.date.today()` is this machine's LOCAL calendar date; requisitions are
# timestamped in UTC (requisitions._now_iso). Without tz_offset_minutes, a
# requisition raised in the gap between local midnight and UTC midnight
# silently falls outside "today" for any timezone east of UTC — exactly
# what this sandbox runs (WAT, UTC+1) and exactly NEEM/TA Connect's own
# timezone. Pass the same offset the frontend now sends
# (Date.getTimezoneOffset()'s own convention) so "today" means the
# server-local calendar day, matching what a NEEM user picking "today" on
# their own machine actually means.
local_offset_minutes = int((dt.datetime.now() - dt.datetime.utcnow()).total_seconds() // 60) * -1

r = client.get("/requisitions/export/log.xlsx", headers=ph,
                params={"start": today_s, "end": tomorrow_s,
                        "tz_offset_minutes": local_offset_minutes})
check("log export (today..tomorrow) succeeds", r.status_code == 200)
wb2 = load_workbook(io.BytesIO(r.content))
log_values = [c.value for row in wb2["Requisition Log"].iter_rows() for c in row]
check("today's requisition IS in the in-range log", req_ref in log_values)

r = client.get("/requisitions/export/log.xlsx", headers=ph,
                params={"start": far_past_start, "end": far_past_end,
                        "tz_offset_minutes": local_offset_minutes})
check("empty range 404s", r.status_code == 404)

# ─── malformed / inverted dates ─────────────────────────────────────────────

r = client.get("/requisitions/export/log.xlsx", headers=ph,
                params={"start": "not-a-date", "end": today_s})
check("malformed start date refused", r.status_code == 400)

r = client.get("/requisitions/export/log.xlsx", headers=ph,
                params={"start": today_s, "end": yesterday_s})
check("end before start refused", r.status_code == 400)

# ─── visibility matches viewing: any signed-in user, no extra role gate ───

client.post("/auth/register", headers=ah, json={
    "email": "viewer@neem.org", "name": "Viewer", "password": "viewer-passphrase",
    "department": "compliance", "role": "viewer"})
vh = hdr("viewer@neem.org", "viewer-passphrase")
r = client.get(f"/requisitions/{req_id}/export.pdf", headers=vh)
check("a plain viewer can export what they can already see", r.status_code == 200)

if _fail:
    print(f"\n{_fail} check(s) FAILED.")
    import sys
    sys.exit(1)
print("\nAll requisition-export checks passed.")
