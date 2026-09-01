"""
Extraction tests — every file format a finance team actually sends.

These exist because file handling failed silently in ways that produced
*plausible wrong numbers* rather than errors, which is the dangerous kind of
bug in a payments system:

  * A photographed receipt decoded as binary and reported a supplier called
    "PNG" with no amount.
  * "SUBTOTAL" matched the "TOTAL" pattern, so a VAT invoice extracted the
    pre-tax figure and would have underpaid the vendor by exactly the tax.
  * A damaged upload and a legitimate scan both reported as "scanned", sending
    a reviewer to re-photograph a file that was never readable in the first
    place.

Run: python test_extraction.py

Fixtures are generated here rather than committed, so the suite is
self-contained. OCR assertions skip cleanly when tesseract isn't installed —
they must not fail the build on a machine that lacks the binary.
"""
from __future__ import annotations

import io
import shutil
import tempfile

_TMP = tempfile.mkdtemp(prefix="docex-extract-")

import store  # noqa: E402

store.set_store(store.JsonFileStore(_TMP))

import fast_extract as fx  # noqa: E402
import field_receipts as fr  # noqa: E402

_passed = 0
_failed = 0
_skipped = 0


def check(label: str, condition: bool) -> None:
    global _passed, _failed
    if condition:
        _passed += 1
        print(f"  ok   {label}")
    else:
        _failed += 1
        print(f"  FAIL {label}")


def skip(label: str, why: str) -> None:
    global _skipped
    _skipped += 1
    print(f"  skip {label} ({why})")


# ─── fixtures ───────────────────────────────────────────────────────────────
# One invoice, expressed the several ways a real supplier might send it.
# SUBTOTAL 375,000 + VAT 28,125 = TOTAL 403,125.

VENDOR = "Sahel Catering Services"
TOTAL = 403125.0
SUBTOTAL = 375000.0


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Voucher"
    for row in [
        ["EDUCATION AS A VACCINE"], ["Payment Voucher"], [],
        ["Vendor:", VENDOR], ["Invoice No:", "INV-2026-0412"],
        ["Date:", "2026-08-14"], ["Project Code:", "P-101"], [],
        ["Description", "Qty", "Unit Price", "Amount"],
        ["Training refreshments", 90, 2500, 225000],
        ["Venue hire", 2, 75000, 150000], [],
        ["", "", "SUBTOTAL", 375000],
        ["", "", "VAT 7.5%", 28125],
        ["", "", "TOTAL", 403125],
    ]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(VENDOR, 0)
    d.add_paragraph("INVOICE INV-2026-0412")
    d.add_paragraph("Vendor: " + VENDOR)
    d.add_paragraph("Date: 2026-08-14")
    t = d.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = "SUBTOTAL"
    t.rows[0].cells[1].text = "375,000.00"
    t.rows[1].cells[0].text = "VAT 7.5%"
    t.rows[1].cells[1].text = "28,125.00"
    t.rows[2].cells[0].text = "TOTAL"
    t.rows[2].cells[1].text = "403,125.00"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def reportlab_available() -> bool:
    """reportlab only builds PDF fixtures for these tests — it is not a runtime
    dependency, so it stays out of requirements.txt and the PDF cases skip
    rather than fail on a machine that doesn't have it."""
    try:
        import reportlab  # noqa: F401

        return True
    except ImportError:
        return False


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in [
        VENDOR.upper(), "Plot 14, Wuse II, Abuja", "", "INVOICE",
        "Invoice No: INV-2026-0412", "Date: 2026-08-14", "",
        "Training refreshments      225,000.00", "Venue hire   150,000.00",
        "SUBTOTAL                   375,000.00", "VAT 7.5%      28,125.00",
        "TOTAL                      403,125.00",
    ]:
        c.drawString(60, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def make_photo() -> bytes:
    """A phone photo of a market receipt — no text layer at all."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (800, 500), "white")
    dr = ImageDraw.Draw(img)
    for i, t in enumerate([
        "MAMA NGOZI PROVISIONS", "DATE: 2026-08-15",
        "Stationery for training", "TOTAL: 47,500.00",
    ]):
        dr.text((40, 60 + i * 60), t, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ─── format coverage ────────────────────────────────────────────────────────


def test_formats_read() -> None:
    print("\nEvery format a finance team sends is readable")
    cases = [
        ("voucher.xlsx", make_xlsx()),
        ("invoice.docx", make_docx()),
        ("receipts.csv", b"vendor,amount\nSahel,403125\n"),
        ("receipt.txt", b"MAMA NGOZI\nTOTAL: 47,500.00\n"),
    ]
    if reportlab_available():
        cases.insert(2, ("invoice.pdf", make_pdf()))
    else:
        skip("invoice.pdf", "reportlab not installed")
    for name, data in cases:
        text, status = fx.extract_status(name, data)
        check(f"{name} extracts text", status == "ok" and len(text.strip()) > 0)

    text, _ = fx.extract_status("voucher.xlsx", make_xlsx())
    check("xlsx keeps the vendor name", VENDOR.lower() in text.lower())
    if reportlab_available():
        text, _ = fx.extract_status("invoice.pdf", make_pdf())
        check("pdf keeps the total", "403,125" in text)


def test_binary_is_never_decoded_as_text() -> None:
    print("\nBinary is never returned as text")
    photo = make_photo()
    # Regardless of OCR, the PNG header must never come back as document text.
    text, status = fx.extract_status("receipt.png", photo)
    check("a photo is not decoded as binary garbage", "IHDR" not in text)
    check("a photo is classified as image, not text", status in ("ocr", "scanned"))

    # A file LYING about its type — a photo named .txt. Sniffing catches it.
    text, status = fx.extract_status("receipt.txt", photo)
    check("magic bytes beat a wrong extension", "IHDR" not in text)


def test_damaged_is_not_called_scanned() -> None:
    print("\nA damaged file is distinguished from a scan")
    text, status = fx.extract_status("broken.pdf", b"%PDF-1.4\n truncated garbage")
    check("damaged PDF reports 'unreadable'", status == "unreadable")
    check("damaged PDF yields no text", text.strip() == "")
    check("damaged PDF is not treated as a scan",
          not fx.is_probably_scanned("broken.pdf", b"%PDF-1.4\n trunc"))


def test_ocr_reads_a_photo() -> None:
    print("\nOCR reads a photographed receipt")
    if not fx.ocr_available():
        skip("photo OCR", "tesseract not installed")
        return
    text, status = fx.extract_status("receipt.png", make_photo())
    check("photo is OCR'd", status == "ocr")
    check("OCR recovers the total", "47,500" in text or "47500" in text)
    check("OCR recovers the vendor", "NGOZI" in text.upper())


# ─── the money bug ──────────────────────────────────────────────────────────


def test_subtotal_is_not_mistaken_for_total() -> None:
    print("\nSUBTOTAL is never read as TOTAL (would underpay by the tax)")
    cases = [("voucher.xlsx", make_xlsx()), ("invoice.docx", make_docx())]
    if reportlab_available():
        cases.insert(0, ("invoice.pdf", make_pdf()))
    else:
        skip("invoice.pdf SUBTOTAL case", "reportlab not installed")
    for name, data in cases:
        r = fr.upload_receipt(
            org_id="t", uploaded_by="fw@t.org", file_content=data, filename=name,
            amount_submitted=TOTAL, project_code="P-101", category="supplies",
        )
        got = r.extracted.extracted_amount
        check(f"{name} reads {TOTAL:,.0f} not {SUBTOTAL:,.0f} (got {got})", got == TOTAL)


def test_vendor_is_clean() -> None:
    print("\nVendor names come through clean")
    r = fr.upload_receipt(
        org_id="t", uploaded_by="fw@t.org", file_content=make_xlsx(),
        filename="voucher.xlsx", amount_submitted=TOTAL,
        project_code="P-101", category="supplies",
    )
    v = r.extracted.vendor_name
    check(f"no leading punctuation (got {v!r})", v == VENDOR)


def test_unreadable_receipt_says_so() -> None:
    print("\nAn unreadable receipt says so, and says which kind")
    r = fr.upload_receipt(
        org_id="t", uploaded_by="fw@t.org",
        file_content=b"%PDF-1.4\n truncated", filename="broken.pdf",
        amount_submitted=1000, project_code="P-101", category="supplies",
    )
    msgs = " ".join(f.message.lower() for f in r.flags)
    check("flagged", r.status == fr.ReceiptStatus.FLAGGED)
    check("says the file is damaged, not that the vendor is missing",
          "damaged" in msgs or "could not be read" in msgs)
    check("does not claim the vendor is simply missing",
          not any(f.type == fr.ReceiptFlagType.MISSING_VENDOR for f in r.flags))


def test_ocr_receipt_is_marked_for_review() -> None:
    print("\nOCR'd figures are marked for human confirmation")
    if not fx.ocr_available():
        skip("OCR review flag", "tesseract not installed")
        return
    r = fr.upload_receipt(
        org_id="t", uploaded_by="fw@t.org", file_content=make_photo(),
        filename="market.png", amount_submitted=47500,
        project_code="P-101", category="supplies",
    )
    check("amount read off the photo", r.extracted.extracted_amount == 47500.0)
    check("confidence stays low for OCR", r.extracted.confidence == "low")
    check("carries a needs_review flag",
          any(f.type == fr.ReceiptFlagType.NEEDS_REVIEW for f in r.flags))


def main() -> int:
    print("=" * 64)
    print("Extraction — formats, OCR, and amounts that must not be wrong")
    print("=" * 64)
    print(f"OCR available: {fx.ocr_available()}")

    test_formats_read()
    test_binary_is_never_decoded_as_text()
    test_damaged_is_not_called_scanned()
    test_ocr_reads_a_photo()
    test_subtotal_is_not_mistaken_for_total()
    test_vendor_is_clean()
    test_unreadable_receipt_says_so()
    test_ocr_receipt_is_marked_for_review()

    print("\n" + "=" * 64)
    print(f"{_passed} passed, {_failed} failed, {_skipped} skipped")
    print("=" * 64)
    return 1 if _failed else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    finally:
        shutil.rmtree(_TMP, ignore_errors=True)
