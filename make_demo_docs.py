#!/usr/bin/env python3
"""
Generate a folder of realistic documents to run through DOCex live.

    python3 make_demo_docs.py

Writes to ./demo_docs/. Everything is fictional.

Why generate files instead of seeding rows: a live demo where you drag a real
invoice in and the system reads it is far more convincing than a screen of
data that was already there. Each file below is chosen to exercise one
specific behaviour, so you can walk the whole product in about ten minutes
without anything being pre-baked.

  01  clean supplier invoice (PDF)     — reads the TOTAL, not the SUBTOTAL
  02  payment voucher (Excel)          — spreadsheets are read too
  03  market receipt (photo)           — OCR; the field-operations story
  04  transport receipt (incomplete)   — flags the gap, invents nothing
  05  duplicate invoice (PDF)          — same invoice number as 01
  06  fuel receipt (scanned PDF)       — no text layer, OCR on a PDF
  07  damaged upload (PDF)             — says "damaged", not "vendor missing"
"""
from __future__ import annotations

import pathlib
import sys

OUT = pathlib.Path(__file__).parent / "demo_docs"


def _canvas(path, lines, size=12):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    y = 790
    for line in lines:
        if line.startswith("##"):
            c.setFont("Helvetica-Bold", size + 3)
            c.drawString(56, y, line[2:].strip())
            c.setFont("Helvetica", size)
        else:
            c.setFont("Helvetica", size)
            c.drawString(56, y, line)
        y -= 19
    c.save()


def invoice_lines(inv_no: str, date: str) -> list[str]:
    return [
        "## KADUNA VENUE SERVICES LTD",
        "12 Independence Way, Kaduna",
        "TIN: 20481937-0001    Tel: 0803 555 0142",
        "",
        "## INVOICE",
        f"Invoice No: {inv_no}",
        f"Date: {date}",
        "Bill To: Community Health Partners",
        "Project: GF-2026-TB",
        "",
        "Description                          Amount (NGN)",
        "-------------------------------------------------",
        "Hall hire, 2 days                       90,000.00",
        "Public address system                   18,000.00",
        "Chairs and tables                       12,000.00",
        "",
        "SUBTOTAL                               120,000.00",
        "VAT 7.5%                                 9,000.00",
        "TOTAL                                  129,000.00",
        "",
        "Payment within 30 days.",
        "Account: 0123456789  Zenith Bank",
    ]


def make_pdf_invoice() -> None:
    _canvas(OUT / "01_invoice_venue_hire.pdf",
            invoice_lines("KVS-2026-0881", "2026-08-18"))


def make_duplicate_invoice() -> None:
    """Same invoice number and amount as 01 — a resubmitted invoice."""
    _canvas(OUT / "05_invoice_venue_hire_RESUBMITTED.pdf",
            invoice_lines("KVS-2026-0881", "2026-08-18"))


def make_voucher_xlsx() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Payment Voucher"
    rows = [
        ["COMMUNITY HEALTH PARTNERS"],
        ["PAYMENT VOUCHER"],
        [],
        ["Vendor:", "Sahel Catering Services"],
        ["Invoice No:", "SCS-2026-0233"],
        ["Date:", "2026-08-20"],
        ["Project Code:", "GF-2026-TB"],
        ["Grant Code:", "GF-2026-TB"],
        ["Activity:", "Community TB screening — Kano, 2 days"],
        [],
        ["Description", "Qty", "Rate", "Amount"],
        ["Participant refreshments, day 1", 45, 2500, 112500],
        ["Participant refreshments, day 2", 45, 2500, 112500],
        ["Facilitator meals", 6, 4000, 24000],
        [],
        ["", "", "SUBTOTAL", 249000],
        ["", "", "VAT 7.5%", 18675],
        ["", "", "TOTAL", 267675],
        [],
        ["Prepared by:", "A. Bello"],
        ["Approved by:", "________________"],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["D"].width = 14
    wb.save(OUT / "02_payment_voucher_training.xlsx")


def _handwritten(lines, width=900, height=620, jitter=True):
    """A photo of a receipt — imperfect on purpose."""
    import random

    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (width, height), (252, 250, 244))
    dr = ImageDraw.Draw(img)

    font = None
    for candidate in (
        "/System/Library/Fonts/Supplemental/Courier New.ttf",
        "/System/Library/Fonts/Menlo.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    ):
        try:
            font = ImageFont.truetype(candidate, 30)
            break
        except Exception:
            continue

    random.seed(7)
    y = 46
    for line in lines:
        x = 46 + (random.randint(-3, 3) if jitter else 0)
        dr.text((x, y), line, fill=(28, 28, 32), font=font)
        y += 52

    # A little grain, so it reads as a photo rather than a screenshot.
    if jitter:
        px = img.load()
        for _ in range(2600):
            gx, gy = random.randint(0, width - 1), random.randint(0, height - 1)
            v = random.randint(205, 240)
            px[gx, gy] = (v, v, v)
    return img


def make_market_photo() -> None:
    img = _handwritten([
        "MAMA NGOZI PROVISIONS",
        "Sabon Gari Market, Kano",
        "",
        "DATE: 2026-08-21",
        "",
        "Exercise books x40",
        "Pens and markers",
        "Flip chart paper",
        "",
        "TOTAL: 47,500.00",
        "",
        "Thank you",
    ])
    img.save(OUT / "03_market_receipt_photo.png")


def make_incomplete_receipt() -> None:
    """No vendor name anywhere — the NEEM pain point, exactly."""
    (OUT / "04_transport_receipt_incomplete.txt").write_text(
        "DATE: 2026-08-22\n"
        "Transport, Kano to Dawakin Kudu (return)\n"
        "3 field officers\n"
        "TOTAL: 18,000.00\n"
    )


def make_scanned_pdf() -> None:
    """An image-only PDF: no text layer at all, OCR is the only way in."""
    import io

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = _handwritten([
        "TOTAL FILLING STATION",
        "Zaria Road, Kano",
        "",
        "DATE: 2026-08-19",
        "",
        "PMS 40 litres",
        "Vehicle: KN-442-ABC",
        "",
        "TOTAL: 62,000.00",
    ], width=840, height=520)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    c = canvas.Canvas(str(OUT / "06_fuel_receipt_scan.pdf"), pagesize=A4)
    c.drawImage(ImageReader(buf), 60, 420, width=470, height=290)
    c.save()


def make_damaged() -> None:
    (OUT / "07_damaged_upload.pdf").write_bytes(
        b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<</Type/Catalog"
        b"\n[truncated during upload]"
    )


def main() -> int:
    OUT.mkdir(exist_ok=True)

    try:
        import reportlab  # noqa: F401
    except ImportError:
        print("error: reportlab is needed to build the PDF samples.\n"
              "       pip3 install reportlab", file=sys.stderr)
        return 1

    make_pdf_invoice()
    make_voucher_xlsx()
    make_market_photo()
    make_incomplete_receipt()
    make_duplicate_invoice()
    make_scanned_pdf()
    make_damaged()

    print(f"Wrote {len(list(OUT.iterdir()))} documents to {OUT}\n")
    for f in sorted(OUT.iterdir()):
        print(f"  {f.name:44} {f.stat().st_size:>8,} bytes")
    print("\nAll fictional. Drag these in during the demo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
